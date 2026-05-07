from openpyxl import Workbook, load_workbook

from app.storage.main import createFolders


def guardar_usuarios_no_encontrados(usuarios_no_encontrados: list[dict]) -> None:
    if not usuarios_no_encontrados:
        return

    ruta_excel = createFolders() / "usuarios_no_encontrados.xlsx"
    encabezados = [
        "fecha",
        "codigo_oficina_planilla",
        "codigo_usuario_volante",
        "usuario_no_encontrado",
        "producto",
        "fecha_activacion",
        "valor_planilla",
        "error_tns",
    ]

    if ruta_excel.exists():
        libro = load_workbook(ruta_excel)
        hoja = libro.active
    else:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Usuarios no encontrados"
        hoja.append(encabezados)

    for usuario in usuarios_no_encontrados:
        hoja.append([
            usuario.get("fecha_volante", ""),
            usuario.get("codigo_oficina_planilla", ""),
            usuario.get("codigo_usuario_volante", ""),
            usuario.get("usuario", ""),
            usuario.get("producto", ""),
            usuario.get("fecha_activacion", ""),
            usuario.get("vr_total_planilla", ""),
            usuario.get("error_tns", ""),
        ])

    libro.save(ruta_excel)
