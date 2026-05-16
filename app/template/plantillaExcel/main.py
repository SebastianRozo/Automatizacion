from openpyxl import load_workbook
from app.template.codes_places.main import OFFICES
from datetime import datetime,date
from app.config.settings import EXCEL_OUTPUT_PATH, EXCEL_TEMPLATE_PATH


def normalizar_texto_excel(valor):
    return str(valor or "").strip().upper()

def obtener_oficina_desde_item(item):
    codigo_usuario = str(item["codigo_usuario"]).strip()
    for oficina in OFFICES:
        usuarios = [str(usuario).strip() for usuario in oficina.get("usuarios", [])]
        if codigo_usuario in usuarios:
            return oficina["name"].upper()
    raise ValueError(f"No se encontro oficina para el usuario: {codigo_usuario}")

def normalizar_fecha_excel(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    valor = str(valor).strip()
    for formato in ("%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(valor, formato).date()
        except ValueError:
            pass
    raise ValueError(f"Fecha no valida: {valor}")
def buscar_fila_por_fecha(hoja, fecha_buscada):
    fecha_buscada = normalizar_fecha_excel(fecha_buscada)
    for fila in range(9, hoja.max_row + 1):
        fecha_excel = hoja.cell(row=fila, column=1).value
        if fecha_excel and normalizar_fecha_excel(fecha_excel) == fecha_buscada:
            return fila
    raise ValueError(f"No se encontro fila para la fecha: {fecha_buscada}")

def buscar_columna_por_oficina(hoja, oficina_buscada):
    oficina_buscada = normalizar_texto_excel(oficina_buscada)
    for columna in range(2, hoja.max_column + 1):
        oficina_excel = normalizar_texto_excel(hoja.cell(row=8, column=columna).value)
        if oficina_excel == oficina_buscada:
            return columna
    raise ValueError(f"No se encontro columna para la oficina: {oficina_buscada}")

def actualizar_excel(datos_volantes: list[dict]):
    try:
        plantilla = load_workbook(EXCEL_TEMPLATE_PATH)
        hoja = plantilla["M-GO-FT-02 V. POSTPAGO"]
        for datos_volante in datos_volantes:
            fecha = datos_volante["fecha"]
            vr_total = datos_volante["dinero"]
            oficina = obtener_oficina_desde_item(datos_volante)
            fila = buscar_fila_por_fecha(hoja, fecha)
            columna = buscar_columna_por_oficina(hoja, oficina)
            valor_actual = hoja.cell(row=fila, column=columna).value or 0
            hoja.cell(row=fila, column=columna, value=valor_actual + vr_total)

        EXCEL_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        plantilla.save(EXCEL_OUTPUT_PATH)
    except Exception as e:
        raise Exception(f"Error al actualizar el excel: {e}") from e
