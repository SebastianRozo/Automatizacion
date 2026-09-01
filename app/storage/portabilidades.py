from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.storage.main import createFolders


NOMBRE_ARCHIVO = "portabilidades_faltantes.xlsx"
NOMBRE_HOJA = "Portabilidades faltantes"
ENCABEZADOS = ["Oficina", "Código", "Cantidad"]


def guardar_portas_faltantes(
    oficina: str,
    codigo: str,
    cantidad: int,
    fecha: str,
) -> Path | None:
    cantidad = int(cantidad or 0)
    if cantidad <= 0:
        return None

    ruta_reporte = createFolders(fecha) / NOMBRE_ARCHIVO
    if ruta_reporte.exists():
        libro = load_workbook(ruta_reporte)
        hoja = libro[NOMBRE_HOJA]
    else:
        libro = Workbook()
        hoja = libro.active
        hoja.title = NOMBRE_HOJA
        hoja.append(ENCABEZADOS)

    codigo = str(codigo).strip()
    fila_oficina = None
    for fila in range(2, hoja.max_row + 1):
        if str(hoja.cell(fila, 2).value or "").strip() == codigo:
            fila_oficina = fila
            break

    if fila_oficina is None:
        hoja.append([oficina, codigo, cantidad])
        fila_oficina = hoja.max_row
    else:
        hoja.cell(fila_oficina, 1, oficina)
        cantidad_actual = int(hoja.cell(fila_oficina, 3).value or 0)
        hoja.cell(fila_oficina, 3, cantidad_actual + cantidad)

    hoja.cell(fila_oficina, 2).number_format = "@"
    libro.save(ruta_reporte)
    libro.close()
    print(f"Portabilidades faltantes guardadas en: {ruta_reporte}")
    return ruta_reporte
